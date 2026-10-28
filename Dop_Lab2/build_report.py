"""Сборка отчёта: титульный лист берётся из отчёта Доп_Lab_1 (тот же титульный лист, что в Lab4), остальное генерируется.
Структура: вариант 1 (задание, текст программы, экранные формы), затем вариант 2."""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

TEMPLATE = "/Users/macbook/Desktop/Доп_Lab_1/Отчет доп_lab_1.docx"
OUT = "Отчет веб-приложение.docx"
TITLE = "по лабораторной работе «Создание веб-приложения для демонстрации моделей машинного обучения»"

doc = Document(TEMPLATE)
body = doc.element.body
paras = doc.paragraphs
idx = next(i for i, p in enumerate(paras) if p.text.startswith("Москва"))
for el in list(body)[list(body).index(paras[idx]._p) + 1:]:
    if not el.tag.endswith("sectPr"):
        body.remove(el)
for p in doc.paragraphs:
    if p.text.startswith("по лабораторной работе"):
        p.runs[0].text = TITLE
        for r in p.runs[1:]:
            r.text = ""

def para(text="", bold=False, size=13, align=WD_ALIGN_PARAGRAPH.JUSTIFY, font=None, after=4):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    r.bold, r.font.size = bold, Pt(size)
    if font:
        r.font.name = font
    return p

def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(0)
    p.add_run(text).font.size = Pt(13)

def code(path, title):
    para(title, bold=True, size=13, align=WD_ALIGN_PARAGRAPH.LEFT, after=2)
    para(open(path, encoding="utf8").read().rstrip(), size=8, align=WD_ALIGN_PARAGRAPH.LEFT, font="Courier New", after=8)

def shots(items, width):
    for name, caption in items:
        para(caption, size=11, align=WD_ALIGN_PARAGRAPH.LEFT, after=2)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(8)
        p.add_run().add_picture(f"screenshots/{name}.png", width=Cm(width))

doc.add_page_break()
para("Описание задания", bold=True, size=18, align=None, after=6)
para("Ссылка на постановку: https://github.com/ugapanyuk/courses_current/wiki/LAB_TMO_WEB", size=13)
para("Цель: изучение возможностей демонстрации моделей машинного обучения с помощью веб-приложений.", size=13)
para("Задание: разработать макет веб-приложения для анализа данных (фреймворк Streamlit).", size=13)
bullet("Вариант 1 — макет для одной модели: задание гиперпараметров, обучение, просмотр результатов (в том числе графиков).")
bullet("Вариант 2 — макет для нескольких моделей: выбор моделей, обучение, просмотр результатов (в том числе графиков).")
para("Датасет: Breast Cancer Wisconsin (Kaggle, uciml/breast-cancer-wisconsin-data), 569 объектов, 30 числовых признаков; "
     "задача — бинарная классификация опухоли (M — злокачественная, B — доброкачественная). "
     "Запуск: streamlit run app_variant1.py и streamlit run app_variant2.py.", size=13)

# ---------------- Вариант 1 ----------------
doc.add_page_break()
para("Вариант 1. Одна модель (случайный лес)", bold=True, size=18, align=None, after=6)
para("Пользователь задаёт гиперпараметры случайного леса (число деревьев, глубину, min_samples_split, max_features), "
     "долю тестовой выборки и random_state, нажимает «Обучить модель» и видит метрики (Accuracy, Precision, Recall, F1, ROC AUC), "
     "матрицу ошибок, ROC-кривую и важность признаков.", size=13)
para("1.1. Текст программы", bold=True, size=14, align=None, after=4)
code("common.py", "common.py (общие функции)")
code("app_variant1.py", "app_variant1.py")
para("1.2. Экранные формы", bold=True, size=14, align=None, after=4)
shots([("v1_01_start", "Начальный экран: гиперпараметры по умолчанию, модель ещё не обучена."),
       ("v1_02_default", "Результаты обучения с параметрами по умолчанию."),
       ("v1_03_tuned", "Изменены гиперпараметры (300 деревьев, max_depth = 8, доля теста 0.30) и модель обучена заново.")], 16)

# ---------------- Вариант 2 ----------------
doc.add_page_break()
para("Вариант 2. Несколько моделей", bold=True, size=18, align=None, after=6)
para("Пользователь выбирает любые из шести моделей (логистическая регрессия, k ближайших соседей, SVM, дерево решений, "
     "случайный лес, градиентный бустинг), задаёт для каждой ключевой гиперпараметр и обучает их на одной и той же выборке. "
     "Приложение показывает таблицу метрик с выделением лучших значений, столбчатую диаграмму, общую ROC-кривую "
     "и матрицы ошибок для каждой модели. Функции загрузки данных и построения графиков — те же, что в common.py.", size=13)
para("2.1. Текст программы", bold=True, size=14, align=None, after=4)
code("app_variant2.py", "app_variant2.py")
para("2.2. Экранные формы", bold=True, size=14, align=None, after=4)
shots([("v2_01_start", "Начальный экран: выбраны две модели по умолчанию."),
       ("v2_02_two_models", "Результаты обучения двух моделей."),
       ("v2_03_all_models", "Выбраны все шесть моделей: таблица метрик, диаграмма, ROC-кривые и матрицы ошибок.")], 16)

para("Выводы", bold=True, size=18, align=None, after=6)
bullet("Разработаны два макета веб-приложений на Streamlit: для одной модели с настройкой гиперпараметров и для сравнения нескольких моделей.")
bullet("Лучшая модель по ROC AUC на тестовой выборке при параметрах по умолчанию — SVM (0.997); дерево решений глубины 4 заметно слабее остальных (Recall 0.736).")
bullet("Результаты зависят от разбиения выборки: тестовая часть содержит около 140 объектов, поэтому различия в 1–2 объекта несущественны.")
doc.save(OUT)
print("saved", OUT)
