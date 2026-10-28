"""Сборка отчёта: титульный лист берётся из отчёта Lab4, остальное генерируется."""
import copy, json
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

TEMPLATE = "/Users/macbook/Desktop/Lab4/Отчет лаба4.docx"
OUT = "Отчет временной ряд.docx"
SHOTS = [2, 3, 5, 8, 9, 10, 11, 12, 13, 14, 15]   # номера ячеек с кодом для экранных форм

doc = Document(TEMPLATE)
body = doc.element.body
paras = doc.paragraphs
idx = next(i for i, p in enumerate(paras) if p.text.startswith("Москва"))
for el in list(body)[list(body).index(paras[idx]._p) + 1:]:
    if el.tag.endswith("sectPr"):
        continue
    body.remove(el)

# титульный лист: номер работы -> название
for p in doc.paragraphs:
    if p.text.startswith("по лабораторной работе"):
        p.runs[0].text = "по лабораторной работе «Анализ и прогнозирование временного ряда»"
        for r in p.runs[1:]:
            r.text = ""

def para(text="", bold=False, size=13, align=WD_ALIGN_PARAGRAPH.JUSTIFY, font=None, after=4):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    if font:
        r.font.name = font
    return p

def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(0)
    p.add_run(text).font.size = Pt(13)

nb = json.load(open("Lab_TimeSeries.ipynb", encoding="utf8"))
code_cells = [c for c in nb["cells"] if c["cell_type"] == "code"]

doc.add_page_break()
para("1. Описание задания", bold=True, size=18, align=None, after=6)
para("Ссылка на постановку: https://github.com/ugapanyuk/courses_current/wiki/LAB_TMO_TIMESERIES", size=13)
para("Цель: изучение основных методов анализа и прогнозирования временных рядов.", size=13)
para("Задание:", bold=True)
for t in ["Выбрать набор данных для задачи прогнозирования временного ряда.",
          "Визуализировать временной ряд и его основные характеристики.",
          "Разделить ряд на обучающую и тестовую выборку.",
          "Выполнить прогноз: авторегрессионным методом (SARIMA), символьной регрессией, "
          "двумя методами МГУА (COMBI и MIA, библиотека gmdh).",
          "Визуализировать тестовую выборку и каждый из прогнозов.",
          "Оценить качество каждого прогноза подходящими метриками."]:
    bullet(t)
para("Датасет: Air Passengers (Kaggle, rakannimer/air-passengers) — ежемесячное число пассажиров авиалиний, "
     "1949–1960, 144 наблюдения. Обучающая выборка — 120 месяцев, тестовая — последние 24.", size=13)

para("2. Текст программы", bold=True, size=18, align=None, after=6)
for c in code_cells:
    p = para(c["source"], size=8, align=WD_ALIGN_PARAGRAPH.LEFT, font="Courier New", after=6)

para("3. Экранные формы", bold=True, size=18, align=None, after=6)
for n in SHOTS:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    p.add_run().add_picture(f"screenshots/cell_{n:02d}.png", width=Cm(16))

para("4. Выводы", bold=True, size=18, align=None, after=6)
for t in ["Ряд имеет тренд и мультипликативную сезонность (период 12), поэтому модели строились для log(y).",
          "Лучшие результаты: МГУА COMBI (RMSE 15.3, MAPE 2.8 %) и символьная регрессия (RMSE 16.9, MAPE 3.0 %); "
          "обе нашли по сути одно правило — значение 12 месяцев назад плюс постоянный прирост в логарифмах.",
          "МГУА MIA (RMSE 36.2) и SARIMA (RMSE 43.1) оказались хуже.",
          "Результаты МГУА сильно зависят от параметра test_size, а тестовая выборка мала (24 точки), "
          "поэтому сравнение методов ориентировочное."]:
    bullet(t)
doc.save(OUT)
print("saved", OUT)
