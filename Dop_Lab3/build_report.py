"""Сборка отчёта: титульный лист берётся из отчёта Доп_Lab_1 (тот же, что в Lab4), остальное генерируется."""
import json
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

TEMPLATE = "/Users/macbook/Desktop/Доп_Lab_1/Отчет доп_lab_1.docx"
OUT = "Отчет кластеризация.docx"
SHOTS = [2, 3, 5, 6, 7, 8, 9, 10, 11, 12]   # номера ячеек с кодом для экранных форм
TITLE = "по лабораторной работе «Обучение без учителя: кластеризация и понижение размерности»"

doc = Document(TEMPLATE)
body = doc.element.body
paras = doc.paragraphs
idx = next(i for i, p in enumerate(paras) if p.text.startswith("Москва"))
for el in list(body)[list(body).index(paras[idx]._p) + 1:]:
    if not el.tag.endswith("sectPr"):
        body.remove(el)
for p in doc.paragraphs:
    if p.text.startswith("по лабораторной работе"):
        p.runs[0].text = TITLE
        for r in p.runs[1:]:
            r.text = ""

def para(text="", bold=False, size=13, align=WD_ALIGN_PARAGRAPH.JUSTIFY, font=None, after=4):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    r.bold, r.font.size = bold, Pt(size)
    if font:
        r.font.name = font
    return p

def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(0)
    p.add_run(text).font.size = Pt(13)

nb = json.load(open("Lab_Cluster.ipynb", encoding="utf8"))
code_cells = [c for c in nb["cells"] if c["cell_type"] == "code"]

doc.add_page_break()
para("1. Описание задания", bold=True, size=18, align=None, after=6)
para("Ссылка на постановку: https://github.com/ugapanyuk/courses_current/wiki/LAB_TMO_CLUSTER", size=13)
para("Цель: изучение методов кластеризации и понижения размерности.", size=13)
para("Задание:", bold=True)
for t in ["Выбрать набор данных для задачи обучения без учителя.",
          "Выбрать подмножество признаков (больше двух, без целевого) — набор D1.",
          "Понизить размерность D1 до двух методом PCA — набор D2.",
          "Понизить размерность D1 до двух методом t-SNE — набор D3.",
          "Визуализировать D2 и D3, определить, где кластеры видны лучше.",
          "Выполнить кластеризацию D1, D2, D3 не менее чем тремя методами и оценить качество метриками из лекционного ноутбука.",
          "Сделать выводы о лучшем методе кластеризации на каждом наборе."]:
    bullet(t)
para("Датасет: Wine (Kaggle, tawfikelmetwally/wine-dataset) — 178 образцов вина трёх сортов, 13 числовых признаков. "
     "Метка сорта в кластеризации не используется, только для оценки (ARI, AMI, гомогенность, полнота, V-мера) и раскраски графиков. "
     "Методы кластеризации: K-Means, Agglomerative (Ward), DBSCAN, GMM.", size=13)

para("2. Текст программы", bold=True, size=18, align=None, after=6)
for c in code_cells:
    para(c["source"], size=8, align=WD_ALIGN_PARAGRAPH.LEFT, font="Courier New", after=6)

para("3. Экранные формы", bold=True, size=18, align=None, after=6)
for n in SHOTS:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    p.add_run().add_picture(f"screenshots/cell_{n:02d}.png", width=Cm(16))

para("4. Выводы", bold=True, size=18, align=None, after=6)
for t in ["На плоскости t-SNE три сорта разделены чётче, чем в PCA (силуэт по истинным меткам 0.550 против 0.522); разница небольшая.",
          "По ARI лучшие методы: K-Means и GMM на D1 (0.880), GMM и K-Means на D2 (0.898 и 0.895), K-Means и GMM на D3 (0.802). "
          "Иерархическая кластеризация хуже (ARI 0.71–0.77).",
          "DBSCAN даёт самый высокий силуэт, но худшие внешние метрики (ARI 0.23–0.40): он относит к шуму 29–42 % точек, "
          "а силуэт шум не учитывает, поэтому «лучший по силуэту» не означает лучший.",
          "Выборка мала (178 образцов), поэтому различия в несколько сотых ARI несущественны."]:
    bullet(t)
doc.save(OUT)
print("saved", OUT)
